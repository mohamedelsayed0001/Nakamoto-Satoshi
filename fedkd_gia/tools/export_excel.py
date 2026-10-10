"""Collect the baseline and every attack experiment under results/ into one Excel workbook.

Sheets:
  Baseline_Rounds   per-round FedKD metrics (no attack)
  Baseline_Clients  per-driver accuracy of the final global student and private teacher
  Baseline_Summary  final and best values (formulas over Baseline_Rounds)
  <exp>             one row per attack (client x setting) + per-setting averages (formulas)
  <exp>_Images      ground truth vs. reconstruction grids
Experiments are folders under results/ containing results.jsonl (searched recursively).

usage: python tools/export_excel.py [--results results] [--out results/fedkd_gia_results.xlsx]
"""
import argparse
import glob
import io
import json
import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="DDEBF7")
SETTING_NAMES = {"none": "S1 no teacher (CE only)", "dinov2": "S2 DINOv2 ViT-B/14 surrogate",
                 "vitb": "S3 unused ImageNet ViT-B/16 surrogate",
                 "none_20k": "S1_full no teacher, 20k iterations",
                 "dinov2_20k": "S2_full DINOv2 surrogate, 20k iterations"}


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(ws, start_row, headers, rows, formats=None):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name=FONT)
            if formats and formats.get(j):
                cell.number_format = formats[j]
    for j in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(40, len(str(headers[j - 1])) + 4))
    return start_row + len(rows)


def baseline_sheets(wb, run_dir, prefix="Baseline", note=None):
    recs = [json.loads(l) for l in open(os.path.join(run_dir, "metrics.jsonl"))]
    ws = wb.create_sheet(f"{prefix}_Rounds")
    headers = ["Round", "Global student acc", "Global student macro-F1", "Mean per-client student acc (held-out drivers for State Farm)",
               "Mean private teacher acc", "Round time (min)", "Mean private teacher macro-F1"]
    rows = [[r["round"], r["global_student_acc"], r["global_student_macro_f1"], r["mean_client_student_acc"],
             r["mean_client_teacher_acc"], r["round_seconds"] / 60, r.get("mean_client_teacher_macro_f1")]
            for r in recs]
    pct = {2: "0.00%", 3: "0.0000", 4: "0.00%", 5: "0.00%", 6: "0.0", 7: "0.0000"}
    last = write_table(ws, 1, headers, rows, pct)
    ws.freeze_panes = "A2"

    wc = wb.create_sheet(f"{prefix}_Clients")
    final = recs[-1]["per_client"]
    rows = [[cid, v.get("student_acc"), v["teacher_test_acc"], v["train_acc_s"], v["train_acc_t"]]
            for cid, v in sorted(final.items())]
    write_table(wc, 1, ["Client", f"Global student test acc (round {recs[-1]['round']})",
                        "Private teacher test acc", "Student train acc (local)", "Teacher train acc (local)"],
                rows, {2: "0.00%", 3: "0.00%", 4: "0.00%", 5: "0.00%"})
    wc.freeze_panes = "A2"
    held = recs[-1].get("per_heldout_driver")
    if held:
        wh = wb.create_sheet(f"{prefix}_HeldOut")
        write_table(wh, 1, ["Held-out driver (never trained on)", "Global student acc (round 30)", "Images"],
                    [[d, v["student_acc"], v["n"]] for d, v in sorted(held.items())], {2: "0.00%"})

    s = wb.create_sheet(f"{prefix}_Summary")
    n = last  # last data row in Baseline_Rounds
    items = [
        ("Final global student acc", f"={prefix}_Rounds!B{n}", "0.00%"),
        ("Final global student macro-F1", f"={prefix}_Rounds!C{n}", "0.0000"),
        ("Final mean private teacher acc", f"={prefix}_Rounds!E{n}", "0.00%"),
        ("Best global student acc", f"=MAX({prefix}_Rounds!B2:B{n})", "0.00%"),
        ("Best global student macro-F1", f"=MAX({prefix}_Rounds!C2:C{n})", "0.0000"),
        ("Best mean private teacher acc", f"=MAX({prefix}_Rounds!E2:E{n})", "0.00%"),
        ("Final mean private teacher macro-F1", f"={prefix}_Rounds!G{n}", "0.0000"),
        ("Total training time (h)", f"=SUM({prefix}_Rounds!F2:F{n})/60", "0.00"),
    ]
    write_table(s, 1, ["Metric", "Value"], [])
    for i, (k, f, fmt) in enumerate(items, 2):
        s.cell(row=i, column=1, value=k).font = Font(name=FONT)
        c = s.cell(row=i, column=2, value=f)
        c.font = Font(name=FONT)
        c.number_format = fmt
    s.column_dimensions["A"].width = 34
    s.cell(row=len(items) + 3, column=1,
           value=note or "Setting: 21 clients (one per driver; p064, p066, p072, p075, p081 held out as unseen test drivers), 30 rounds x 5 local epochs, batch 8, ViT-B mentor / ViT-S "
                 "mentee, full FedKD loss, no SVD compression. Accuracy and F1 are measured on the held-out drivers.").font = Font(name=FONT, italic=True)


def experiment_sheets(wb, name, files):
    recs = [dict(json.loads(l), run=os.path.basename(os.path.dirname(f))) for f in files for l in open(f)]
    recs.sort(key=lambda r: (r["run"], list(SETTING_NAMES).index(r["setting"]) if r["setting"] in SETTING_NAMES else 9, r["client"]))
    cfgs = {}
    for f in files:
        cp = os.path.join(os.path.dirname(f), "config.json")
        if os.path.exists(cp):
            c = json.load(open(cp))
            cfgs[os.path.basename(os.path.dirname(f))] = {k: c[k] for k in ("iterations", "num_seeds", "lr", "alpha_tv",
                                                                          "alpha_l2", "alpha_group", "alpha_noise") if k in c}

    ws = wb.create_sheet(name[:31])
    headers = ["Setting", "Client", "Run (config folder)", "Label restoration acc", "PSNR consensus (dB)", "PSNR best seed (dB)",
               "PSNR gray floor (dB)", "Consensus - floor (dB)", "Final grad-match loss (best seed)", "Time (min)"]
    rows = []
    for i, r in enumerate(recs, 2):
        rows.append([SETTING_NAMES.get(r["setting"], r["setting"]), r["client"], r["run"], r["label_acc"],
                     r["psnr_consensus"], r["psnr_best_seed"], r.get("psnr_gray_floor"),
                     f"=E{i}-G{i}" if "psnr_gray_floor" in r else None, min(r["final_grad_loss"]), r["seconds"] / 60])
    last = write_table(ws, 1, headers, rows, {4: "0.0%", 5: "0.00", 6: "0.00", 7: "0.00", 8: "0.00", 9: "0.0000",
                                               10: "0.0"})
    ws.column_dimensions["A"].width = 34
    ws.freeze_panes = "A2"

    # Per-setting averages.
    r0 = last + 3
    ws.cell(row=r0 - 1, column=1, value="Averages per setting").font = Font(name=FONT, bold=True)
    sum_headers = ["Setting", "Attacks", "Label restoration acc", "PSNR consensus (dB)", "PSNR best seed (dB)",
                   "PSNR gray floor (dB)", "Consensus - floor (dB)"]
    write_table(ws, r0, sum_headers, [])
    settings = [s for s in SETTING_NAMES if any(r["setting"] == s for r in recs)]
    rng = lambda col: f"{col}$2:{col}${last}"  # noqa: E731
    for k, s in enumerate(settings, r0 + 1):
        label = SETTING_NAMES[s]
        ws.cell(row=k, column=1, value=label).font = Font(name=FONT)
        ws.cell(row=k, column=2, value=f'=COUNTIF({rng("A")},A{k})')
        for col_idx, src in zip(range(3, 8), "DEFGH"):
            c = ws.cell(row=k, column=col_idx, value=f'=AVERAGEIFS({rng(src)},{rng("A")},A{k})')
            c.number_format = "0.0%" if src == "D" else "0.00"
            c.font = Font(name=FONT)

    note_row = r0 + len(settings) + 2
    notes = [
        "Threat model: honest-but-curious server sees one client's student gradient for one local step "
        "(batch 8, distinct labels, round-30 global student). Private teacher and projector are unknown to the attacker.",
        "PSNR: images in [0,1], reconstructions matched one-to-one to ground truth (Hungarian), averaged over the batch. "
        "Consensus = pixel mean over seeds; best seed = lowest gradient-matching loss (attacker-selectable).",
        "Gray floor: PSNR of a uniform gray image; a reconstruction below it carries no recoverable pixel information.",
    ] + [f"Run {k}: {json.dumps(v)}" for k, v in sorted(cfgs.items())]
    for i, t in enumerate(notes):
        ws.cell(row=note_row + i, column=1, value=t).font = Font(name=FONT, italic=True)

    # Images.
    pngs = sorted(p for f in files for p in glob.glob(os.path.join(os.path.dirname(f), "*.png")))
    if pngs:
        wi = wb.create_sheet((name + "_Images")[:31])
        wi.cell(row=1, column=1, value="Each grid: row 1 = actual batch, row 2 = consensus reconstruction, "
                                       "row 3 = best-seed reconstruction").font = Font(name=FONT, bold=True)
        row = 3
        for p in pngs:
            stem = os.path.basename(p)[:-4]
            setting, client = stem.split("_", 1)
            run = os.path.basename(os.path.dirname(p))
            wi.cell(row=row, column=1, value=f"{SETTING_NAMES.get(setting, setting)} - client {client} - run {run}").font = Font(name=FONT, bold=True)
            # Embed a downscaled JPEG copy so the workbook stays small; full-size PNGs stay in results/.
            im = Image.open(p).convert("RGB")
            w, h = im.size
            small = im.resize((1100, int(1100 * h / w)), Image.LANCZOS)
            buf = io.BytesIO()
            small.save(buf, format="JPEG", quality=85)
            buf.seek(0)
            img = XLImage(buf)
            img.width, img.height = small.size
            wi.add_image(img, f"A{row + 1}")
            row += int(img.height / 20) + 4


def comparison_sheet(wb, results_dir, prefix="gi_s"):
    """Client x setting comparison over the full attack runs (folders starting with `prefix`)."""
    res = {}
    for exp in [e for e in sorted(os.listdir(results_dir)) if not e.startswith("old_")]:
        if not (exp.startswith(prefix) or exp in ("S1_full", "S2_full")):
            continue
        for f in glob.glob(os.path.join(results_dir, exp, "**", "results.jsonl"), recursive=True):
            for l in open(f):
                r = json.loads(l)
                key = {"S1_full": "none_20k", "S2_full": "dinov2_20k"}.get(exp, r["setting"])
                res.setdefault(r["client"], {})[key] = r
    if not res:
        return
    settings = [s for s in SETTING_NAMES if any(s in v for v in res.values())]
    ws = wb.create_sheet("GI_Comparison", 0)
    headers = ["Client", "PSNR gray floor (dB)"]
    for s in settings:
        headers += [f"{s}: PSNR consensus (dB)", f"{s}: gain over floor (dB)", f"{s}: PSNR best seed (dB)"]
    clients = sorted(res)
    rows = []
    for i, c in enumerate(clients, 2):
        floor = next(iter(res[c].values()))["psnr_gray_floor"]
        row = [c, floor]
        for j, s in enumerate(settings):
            col = get_column_letter(3 + 3 * j)
            r = res[c].get(s)
            row += [r["psnr_consensus"] if r else None, f"={col}{i}-$B{i}" if r else None,
                    r["psnr_best_seed"] if r else None]
        rows.append(row)
    last = write_table(ws, 1, headers, rows, {k: "0.00" for k in range(2, len(headers) + 1)})
    mean_row, above_row = last + 1, last + 2
    ws.cell(row=mean_row, column=1, value="Mean").font = Font(name=FONT, bold=True)
    ws.cell(row=above_row, column=1, value="Clients above floor").font = Font(name=FONT, bold=True)
    for k in range(2, len(headers) + 1):
        col = get_column_letter(k)
        c = ws.cell(row=mean_row, column=k, value=f"=AVERAGE({col}2:{col}{last})")
        c.number_format, c.font = "0.00", Font(name=FONT, bold=True)
    for j in range(len(settings)):
        col = get_column_letter(4 + 3 * j)
        ws.cell(row=above_row, column=4 + 3 * j, value=f'=COUNTIF({col}2:{col}{last},">0")').font = Font(name=FONT, bold=True)
    notes = ["Settings: " + "; ".join(f"{s} = {SETTING_NAMES[s]}" for s in settings),
             "Gain over floor = consensus PSNR minus PSNR of a uniform gray image (no-information reference).",
             "All attacks: round-30 victim gradient, batch 8, 2 seeds, TV 1, no Langevin noise; 4000 iterations "
             "except the _20k columns (S1_full / S2_full: 20000 iterations, 6 clients)."]
    for i, t in enumerate(notes, above_row + 2):
        ws.cell(row=i, column=1, value=t).font = Font(name=FONT, italic=True)
    ws.freeze_panes = "B2"


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--results", default=os.path.join(os.path.dirname(os.path.dirname(__file__)), "results"))
    p.add_argument("--out", default=None)
    args = p.parse_args()
    out = args.out or os.path.join(args.results, "fedkd_gia_results.xlsx")

    wb = Workbook()
    wb.remove(wb.active)
    baseline_sheets(wb, os.path.join(args.results, "baseline", "run"))
    cifar = os.path.join(args.results, "cifar100_baseline", "run")
    if os.path.exists(os.path.join(cifar, "metrics.jsonl")):
        baseline_sheets(wb, cifar, prefix="CIFAR100_Baseline",
                        note="Setting: CIFAR-100, 30 IID clients (80/20 local train/test, 1,333 / 334 images each), "
                             "30 rounds x 1 local epoch, batch 8, 32x32 images upsampled to 224, ViT-B mentor / ViT-S "
                             "mentee, full FedKD loss, no SVD. Global student acc/F1: official 10k test set; "
                             "per-client and teacher numbers: each client's local 20% test split.")
    for exp in [e for e in sorted(os.listdir(args.results)) if not e.startswith("old_")]:
        files = sorted(glob.glob(os.path.join(args.results, exp, "**", "results.jsonl"), recursive=True))
        if files:
            experiment_sheets(wb, exp, files)
    comparison_sheet(wb, args.results)
    for exp in [e for e in sorted(os.listdir(args.results)) if not e.startswith("old_")]:
        plots = sorted(glob.glob(os.path.join(args.results, exp, "plots", "*.png")))
        if plots:
            ws = wb.create_sheet((exp + "_Plots")[:31])
            row = 1
            for p in plots:
                ws.cell(row=row, column=1, value=os.path.basename(p)).font = Font(name=FONT, bold=True)
                img = XLImage(p)
                scale = 900 / img.width
                img.width, img.height = 900, int(img.height * scale)
                ws.add_image(img, f"A{row + 1}")
                row += int(img.height / 20) + 4
    wb.save(out)
    print("saved", out, "sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
